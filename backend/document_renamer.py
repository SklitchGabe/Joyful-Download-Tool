import re
import PyPDF2
from pathlib import Path
from typing import Optional, Tuple
import docx

def extract_project_id_from_docx(docx_path: str) -> Optional[str]:
    """
    Extract the first occurrence of a World Bank project ID from a DOCX file.
    Project IDs are in the format P followed by 6 digits (e.g., P123456).
    
    Args:
        docx_path: Path to the DOCX file
        
    Returns:
        The project ID if found, None otherwise
    """
    try:
        # Open the DOCX file
        doc = docx.Document(docx_path)
        
        # Regular expression pattern for project ID
        pattern = r'P\d{6}'
        
        # Extract text from paragraphs
        for para in doc.paragraphs:
            text = para.text
            matches = re.findall(pattern, text)
            if matches:
                return matches[0]
                
        # If not found in paragraphs, try tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    matches = re.findall(pattern, cell.text)
                    if matches:
                        return matches[0]
        
        return None
        
    except Exception as e:
        print(f"Error processing DOCX {docx_path}: {str(e)}")
        return None

def extract_project_id(pdf_path: str, max_pages: int = 10) -> Optional[str]:
    """
    Extract the first occurrence of a World Bank project ID from a PDF file.
    Project IDs are in the format P followed by 6 digits (e.g., P123456).
    
    Args:
        pdf_path: Path to the PDF file
        max_pages: Maximum number of pages to search (default: 10)
        
    Returns:
        The project ID if found, None otherwise
    """
    try:
        with open(pdf_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            # Limit the number of pages to search
            pages_to_search = min(len(reader.pages), max_pages)
            
            # Regular expression pattern for project ID
            pattern = r'P\d{6}'
            
            # Search through pages
            for page_num in range(pages_to_search):
                page = reader.pages[page_num]
                text = page.extract_text()
                
                # Find all matches in the page
                matches = re.findall(pattern, text)
                if matches:
                    return matches[0]  # Return the first match
                    
        return None
        
    except Exception as e:
        print(f"Error processing {pdf_path}: {str(e)}")
        return None

def rename_document_with_project_id(original_path: str, output_dir: str) -> Optional[str]:
    """
    Rename a document file based on the project ID extracted from its content.
    
    Args:
        original_path: Path to the original file
        output_dir: Directory to place the renamed file
        
    Returns:
        Path to the renamed file if successful, original path if no ID found
    """
    try:
        # Get file extension from original path
        file_extension = Path(original_path).suffix.lower()
        
        # Extract project ID from the document based on file type
        project_id = None
        if file_extension == '.pdf':
            project_id = extract_project_id(original_path)
        elif file_extension == '.docx':
            project_id = extract_project_id_from_docx(original_path)
        else:
            # For non-PDF/DOCX files, try to extract from filename
            import re
            match = re.search(r'(P\d{6})', Path(original_path).stem)
            if match:
                project_id = match.group(1)
        
        if project_id:
            # Create the output directory if it doesn't exist
            Path(output_dir).mkdir(parents=True, exist_ok=True)
            
            # Create new filename with project ID
            new_filename = f"{project_id}{file_extension}"
            new_path = Path(output_dir) / new_filename
            
            # Handle duplicate filenames
            counter = 1
            while new_path.exists():
                new_filename = f"{project_id}_{counter}{file_extension}"
                new_path = Path(output_dir) / new_filename
                counter += 1
            
            # Copy the file to the new location with the new name
            import shutil
            shutil.copy2(original_path, new_path)
            
            return str(new_path)
        else:
            # If no project ID found, return the original path
            return original_path
            
    except Exception as e:
        print(f"Error renaming document {original_path}: {str(e)}")
        return original_path